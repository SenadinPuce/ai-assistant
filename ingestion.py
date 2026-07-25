import os
import time
from pathlib import Path
from typing import Iterable

from dotenv import load_dotenv
from langchain_community.document_loaders import PyPDFDirectoryLoader
from langchain_core.documents import Document
from langchain_pinecone import PineconeVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pinecone import Pinecone, ServerlessSpec

load_dotenv()

from retrieval.config import (
    EMBEDDING_DIMENSIONS,
    PDF_DIR,
    PINECONE_INDEX_NAME,
)
from retrieval.vectorstore import _get_embeddings

SUPPORTED_UPLOAD_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".csv",
    ".json",
    ".html",
    ".htm",
}


def _ensure_index_exists() -> None:
    """Create the Pinecone index if missing and wait until it is ready."""
    pc = Pinecone(api_key=os.environ["PINECONE_API_KEY"])

    list_result = pc.list_indexes()
    if hasattr(list_result, "names"):
        index_names = set(list_result.names())
    else:
        index_names = {idx.name for idx in list_result}

    if PINECONE_INDEX_NAME not in index_names:
        pc.create_index(
            name=PINECONE_INDEX_NAME,
            dimension=EMBEDDING_DIMENSIONS,
            metric="cosine",
            spec=ServerlessSpec(cloud="aws", region="us-east-1"),
        )

    max_wait_seconds = 90
    deadline = time.time() + max_wait_seconds
    while time.time() < deadline:
        description = pc.describe_index(PINECONE_INDEX_NAME)
        status = getattr(description, "status", {})
        if isinstance(status, dict) and status.get("ready"):
            return
        if hasattr(status, "ready") and status.ready:
            return
        time.sleep(2)

    raise RuntimeError(
        f"Pinecone index '{PINECONE_INDEX_NAME}' is not ready after {max_wait_seconds}s. "
        "Verify PINECONE_API_KEY/PINECONE_INDEX_NAME point to the same Pinecone project."
    )


def extract_text_from_file(file_path: str | Path) -> str:
    """Extract text from a supported file type."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".csv", ".json", ".html", ".htm"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover - exercised only if package is missing
            raise RuntimeError("python-docx is required to parse .docx files") from exc

        document = DocxDocument(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

    raise ValueError(f"Unsupported file type: {path.suffix or 'unknown'}")


def _load_documents_from_paths(file_paths: Iterable[str | Path]) -> list[Document]:
    documents: list[Document] = []
    for raw_path in file_paths:
        path = Path(raw_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        text = extract_text_from_file(path)
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": path.suffix.lower(),
                },
            )
        )
    return documents


def ingest_documents(
    pdf_dir: str | None = None,
    files: Iterable[str | Path] | None = None,
) -> int:
    """Load documents, split into chunks, and upsert into Pinecone.

    Run this once (or whenever the document set changes).
    """
    _ensure_index_exists()

    if files is not None:
        docs = _load_documents_from_paths(files)
    else:
        docs = PyPDFDirectoryLoader(pdf_dir or PDF_DIR).load()

    doc_splits = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        chunk_size=512, chunk_overlap=64
    ).split_documents(docs)

    PineconeVectorStore.from_documents(
        documents=doc_splits,
        embedding=_get_embeddings(),
        index_name=PINECONE_INDEX_NAME,
    )
    print(f"Ingested {len(doc_splits)} chunks into '{PINECONE_INDEX_NAME}'.")
    return len(doc_splits)


if __name__ == "__main__":
    ingest_documents()
